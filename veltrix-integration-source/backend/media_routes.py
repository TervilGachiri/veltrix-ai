import base64
import io
import os
import re
from pathlib import Path

import httpx
from fastapi import (
    APIRouter, Cookie, File, HTTPException,
    Request, UploadFile
)
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field

from auth_routes import require_user

router = APIRouter(prefix="/api/media", tags=["Media"])

MAX_BYTES = 8 * 1024 * 1024
MAX_TEXT = 18000

ALLOWED = {
    ".pdf", ".docx", ".xlsx", ".txt",
    ".csv", ".png", ".jpg", ".jpeg", ".webp"
}
IMAGE_TYPES = {".png", ".jpg", ".jpeg", ".webp"}


def authenticate(token):
    return require_user(token)


def verify_origin(request: Request):
    origin = request.headers.get("origin")
    allowed = {
        "http://localhost:5173",
        "http://127.0.0.1:5173"
    }
    configured = os.getenv(
        "VELTRIX_FRONTEND_ORIGIN", ""
    ).rstrip("/")
    if configured:
        allowed.add(configured)

    if origin and origin.rstrip("/") not in allowed:
        raise HTTPException(403, "Untrusted request origin.")


async def read_limited(upload: UploadFile):
    data = await upload.read(MAX_BYTES + 1)
    if len(data) > MAX_BYTES:
        raise HTTPException(413, "Maximum file size is 8 MB.")
    if not data:
        raise HTTPException(400, "The file is empty.")
    return data


def file_extension(filename):
    ext = Path(filename or "").suffix.lower()
    if ext not in ALLOWED:
        raise HTTPException(
            415, "Unsupported file format."
        )
    return ext


def extract_content(data, ext):
    if ext in (".txt", ".csv"):
        return data.decode("utf-8-sig", errors="replace")[:MAX_TEXT]

    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        if len(reader.pages) > 150:
            raise HTTPException(413, "Too many PDF pages.")

        parts = []
        total = 0
        for page in reader.pages:
            text = page.extract_text() or ""
            remaining = MAX_TEXT - total
            if remaining <= 0:
                break
            parts.append(text[:remaining])
            total += len(parts[-1])
        return "\n".join(parts)

    if ext == ".docx":
        from docx import Document

        document = Document(io.BytesIO(data))
        output = []
        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                output.append(paragraph.text)
        for table in document.tables:
            for row in table.rows:
                output.append(
                    " | ".join(cell.text for cell in row.cells)
                )
        return "\n".join(output)[:MAX_TEXT]

    if ext == ".xlsx":
        from openpyxl import load_workbook

        book = load_workbook(
            io.BytesIO(data),
            read_only=True,
            data_only=True
        )
        output = []
        remaining = MAX_TEXT

        try:
            for sheet in book.worksheets[:10]:
                output.append(f"\nSheet: {sheet.title}")
                for index, row in enumerate(
                    sheet.iter_rows(values_only=True)
                ):
                    if index >= 300 or remaining <= 0:
                        break
                    line = " | ".join(
                        "" if cell is None else str(cell)
                        for cell in row[:30]
                    )
                    line = line[:remaining]
                    output.append(line)
                    remaining -= len(line)
        finally:
            book.close()

        return "\n".join(output)[:MAX_TEXT]

    raise HTTPException(415, "No text extractor available.")


def validated_image(data):
    from PIL import Image, ImageOps, UnidentifiedImageError

    try:
        image = Image.open(io.BytesIO(data))
        image.verify()

        image = Image.open(io.BytesIO(data))
        if image.width * image.height > 20_000_000:
            raise HTTPException(413, "Image resolution too large.")

        image = ImageOps.exif_transpose(image)
        image.thumbnail((1280, 1280))
        image = image.convert("RGB")

        result = io.BytesIO()
        image.save(result, format="JPEG", quality=82)
        return result.getvalue()
    except HTTPException:
        raise
    except (UnidentifiedImageError, OSError, ValueError):
        raise HTTPException(415, "Invalid or unsupported image.")


@router.post("/extract")
async def extract(
    request: Request,
    file: UploadFile = File(...),
    veltrix_session: str | None = Cookie(default=None),
):
    authenticate(veltrix_session)
    verify_origin(request)

    ext = file_extension(file.filename)
    if ext in IMAGE_TYPES:
        raise HTTPException(
            400,
            "Images need the vision model, not text extraction."
        )

    data = await read_limited(file)

    try:
        content = extract_content(data, ext)
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(
            422,
            "The document could not be read."
        )

    if not content.strip():
        raise HTTPException(
            422,
            "No readable text found. Scanned PDFs require OCR."
        )

    return {
        "filename": Path(file.filename or "document").name,
        "text": content,
        "truncated": len(content) >= MAX_TEXT
    }


@router.post("/vision")
async def vision(
    request: Request,
    file: UploadFile = File(...),
    veltrix_session: str | None = Cookie(default=None),
):
    authenticate(veltrix_session)
    verify_origin(request)

    ext = file_extension(file.filename)
    if ext not in IMAGE_TYPES:
        raise HTTPException(415, "Upload an image.")

    content = await read_limited(file)
    image_bytes = validated_image(content)

    model = os.getenv("OLLAMA_VISION_MODEL", "gemma3:4b")
    payload = {
        "model": model,
        "messages": [{
            "role": "user",
            "content": (
                "Describe this image accurately. "
                "Extract clearly readable text where relevant. "
                "Do not invent details you cannot see."
            ),
            "images": [
                base64.b64encode(image_bytes).decode("ascii")
            ]
        }],
        "stream": False,
        "options": {"num_predict": 400}
    }

    try:
        async with httpx.AsyncClient(
            timeout=240
        ) as client:
            response = await client.post(
                "http://127.0.0.1:11434/api/chat",
                json=payload
            )

        if response.status_code >= 400:
            raise HTTPException(
                503,
                "Vision model unavailable. "
                "Install the configured Ollama vision model."
            )

        result = response.json()
        description = result.get(
            "message", {}
        ).get("content", "").strip()

        description = re.sub(
            r"<think>.*?</think>",
            "",
            description,
            flags=re.DOTALL
        ).strip()

        if "</think>" in description:
            description = description.rsplit(
                "</think>", 1
            )[-1].strip()

        if not description:
            raise HTTPException(
                502, "Vision model returned no description."
            )

        return {
            "filename": Path(file.filename or "photo").name,
            "description": description[:MAX_TEXT]
        }

    except httpx.TimeoutException:
        raise HTTPException(504, "Image analysis timed out.")
    except httpx.RequestError:
        raise HTTPException(
            503, "Ollama vision service is unavailable."
        )


class ExportRequest(BaseModel):
    text: str = Field(min_length=1, max_length=30000)
    format: str = Field(pattern="^(pdf|docx|xlsx|txt)$")


@router.post("/export")
def export(
    data: ExportRequest,
    request: Request,
    veltrix_session: str | None = Cookie(default=None),
):
    authenticate(veltrix_session)
    verify_origin(request)

    content = io.BytesIO()
    fmt = data.format

    if fmt == "pdf":
        from reportlab.platypus import (
            SimpleDocTemplate, Paragraph, Spacer
        )
        from reportlab.lib.styles import getSampleStyleSheet
        from xml.sax.saxutils import escape

        doc = SimpleDocTemplate(content)
        styles = getSampleStyleSheet()

        story = [
            Paragraph("Veltrix AI Report", styles["Title"]),
            Spacer(1, 18)
        ]

        for line in data.text.splitlines():
            if line.strip():
                story.append(
                    Paragraph(
                        escape(line[:2000]),
                        styles["BodyText"]
                    )
                )
                story.append(Spacer(1, 6))

        doc.build(story)
        mime = "application/pdf"

    elif fmt == "docx":
        from docx import Document

        document = Document()
        document.add_heading("Veltrix AI Report", 0)

        for line in data.text.splitlines():
            document.add_paragraph(line)

        document.save(content)
        mime = (
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        )

    elif fmt == "xlsx":
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Veltrix Report"

        for line in data.text.splitlines()[:5000]:
            # Prevent spreadsheet formula injection.
            if line.startswith(("=", "+", "-", "@")):
                line = "'" + line
            sheet.append([line[:32000]])

        workbook.save(content)
        mime = (
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )

    else:
        content.write(data.text.encode("utf-8"))
        mime = "text/plain; charset=utf-8"

    content.seek(0)

    return StreamingResponse(
        content,
        media_type=mime,
        headers={
            "Content-Disposition":
                f'attachment; filename="veltrix-report.{fmt}"',
            "Cache-Control": "no-store"
        }
    )
